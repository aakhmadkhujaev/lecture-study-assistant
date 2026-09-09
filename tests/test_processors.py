"""Unit tests for PDF, DOCX, and PPTX extraction."""

from pathlib import Path

import pymupdf
import pytest
from docx import Document as WordDocument
from pptx import Presentation

from app.processors.docx_processor import DOCXProcessor
from app.processors.factory import extract_document, get_processor
from app.processors.models import (
    DocumentNotFoundError,
    EmptyDocumentError,
    UnsupportedDocumentError,
)
from app.processors.pdf_processor import PDFProcessor
from app.processors.pptx_processor import PPTXProcessor


def _create_pdf(path: Path) -> None:
    document = pymupdf.open()
    for text in ("Page one text", "Page two text"):
        page = document.new_page()
        page.insert_text((72, 72), text)
    document.save(path)
    document.close()


def _create_docx(path: Path) -> None:
    document = WordDocument()
    document.add_heading("Document heading", level=1)
    document.add_paragraph("First paragraph")
    document.add_paragraph("Second paragraph")
    document.save(path)


def _create_pptx(path: Path) -> None:
    presentation = Presentation()
    for title, body in (
        ("Slide one", "First slide text"),
        ("Slide two", "Second slide text"),
    ):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    presentation.save(path)


def test_pdf_extracts_pages_with_traceability(tmp_path: Path) -> None:
    path = tmp_path / "lecture.pdf"
    _create_pdf(path)

    document = PDFProcessor().process(path)

    assert document.filename == "lecture.pdf"
    assert document.file_type == ".pdf"
    assert [(section.source_type, section.source_index) for section in document.sections] == [
        ("page", 1),
        ("page", 2),
    ]
    assert [section.text for section in document.sections] == [
        "Page one text",
        "Page two text",
    ]


def test_docx_preserves_paragraph_order_and_headings(tmp_path: Path) -> None:
    path = tmp_path / "lecture.docx"
    _create_docx(path)

    document = DOCXProcessor().process(path)

    assert [section.text for section in document.sections] == [
        "Document heading",
        "First paragraph",
        "Second paragraph",
    ]
    assert [section.source_type for section in document.sections] == [
        "heading",
        "paragraph",
        "paragraph",
    ]
    assert [section.source_index for section in document.sections] == [1, 2, 3]


def test_pptx_extracts_slides_with_traceability(tmp_path: Path) -> None:
    path = tmp_path / "lecture.pptx"
    _create_pptx(path)

    document = PPTXProcessor().process(path)

    assert [(section.source_type, section.source_index) for section in document.sections] == [
        ("slide", 1),
        ("slide", 2),
    ]
    assert "Slide one" in document.sections[0].text
    assert "First slide text" in document.sections[0].text
    assert "Second slide text" in document.sections[1].text


@pytest.mark.parametrize("extension, processor_type", [
    (".pdf", PDFProcessor),
    (".docx", DOCXProcessor),
    (".pptx", PPTXProcessor),
])
def test_processor_selection(extension: str, processor_type: type, tmp_path: Path) -> None:
    processor = get_processor(tmp_path / f"material{extension.upper()}")
    assert isinstance(processor, processor_type)


def test_unsupported_extension_is_rejected(tmp_path: Path) -> None:
    with pytest.raises(UnsupportedDocumentError):
        get_processor(tmp_path / "notes.txt")


def test_missing_file_is_rejected(tmp_path: Path) -> None:
    with pytest.raises(DocumentNotFoundError):
        extract_document(tmp_path / "missing.pdf")


def test_empty_pdf_is_reported_as_non_extractable(tmp_path: Path) -> None:
    path = tmp_path / "empty.pdf"
    document = pymupdf.open()
    document.new_page()
    document.save(path)
    document.close()

    with pytest.raises(EmptyDocumentError, match="OCR could not detect readable content"):
        PDFProcessor().process(path)


def test_empty_docx_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "empty.docx"
    WordDocument().save(path)

    with pytest.raises(EmptyDocumentError):
        DOCXProcessor().process(path)
