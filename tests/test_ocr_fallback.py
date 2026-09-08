"""Tests for local OCR fallback in PDF and PPTX processors."""

from io import BytesIO
from pathlib import Path

import pymupdf
import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from app.processors.docx_processor import DOCXProcessor
from app.processors.models import EmptyDocumentError
from app.processors.ocr import OCRProcessingError
from app.processors.pdf_processor import PDFProcessor
from app.processors.pptx_processor import PPTXProcessor


class FakeOCR:
    def __init__(self, responses: list[str] | None = None, error: Exception | None = None) -> None:
        self.responses = responses or []
        self.error = error
        self.calls = 0

    def extract_text(self, image: Image.Image) -> str:
        self.calls += 1
        if self.error:
            raise self.error
        return self.responses.pop(0)


def _image_bytes(text: str) -> bytes:
    image = Image.new("RGB", (600, 200), "white")
    drawer = ImageDraw.Draw(image)
    drawer.text((30, 80), text, fill="black")
    output = BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def _create_image_pptx(path: Path, texts: list[str]) -> None:
    presentation = Presentation()
    for text in texts:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(BytesIO(_image_bytes(text)), Inches(1), Inches(1), width=Inches(7))
    presentation.save(path)


def _create_mixed_pdf(path: Path) -> None:
    document = pymupdf.open()
    text_page = document.new_page()
    text_page.insert_text((72, 72), "Normal page text")
    image_page = document.new_page()
    image_page.insert_image(pymupdf.Rect(72, 72, 500, 220), stream=_image_bytes("Scanned page"))
    document.save(path)
    document.close()


def _create_image_pdf(path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_image(pymupdf.Rect(72, 72, 500, 220), stream=_image_bytes("Scanned page"))
    document.save(path)
    document.close()


def test_text_pdf_does_not_call_ocr(tmp_path: Path) -> None:
    path = tmp_path / "text.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Normal page text")
    document.save(path)
    document.close()
    ocr = FakeOCR(["should not be used"])

    result = PDFProcessor(ocr).process(path)

    assert result.extraction_method == "text"
    assert result.sections[0].extraction_method == "text"
    assert ocr.calls == 0


def test_image_pdf_uses_ocr_and_preserves_page_index(tmp_path: Path) -> None:
    path = tmp_path / "scanned.pdf"
    _create_image_pdf(path)
    ocr = FakeOCR(["OCR page text"])

    result = PDFProcessor(ocr).process(path)

    assert result.extraction_method == "ocr"
    assert result.filename == "scanned.pdf"
    assert result.sections[0].source_type == "page"
    assert result.sections[0].source_index == 1
    assert result.sections[0].extraction_method == "ocr"
    assert result.sections[0].text == "OCR page text"
    assert ocr.calls == 1


def test_mixed_pdf_only_ocr_fallback_page_is_marked(tmp_path: Path) -> None:
    path = tmp_path / "mixed.pdf"
    _create_mixed_pdf(path)
    ocr = FakeOCR(["OCR page text"])

    result = PDFProcessor(ocr).process(path)

    assert [section.extraction_method for section in result.sections] == ["text", "ocr"]
    assert [section.source_index for section in result.sections] == [1, 2]


def test_image_pptx_uses_ocr_per_slide(tmp_path: Path) -> None:
    path = tmp_path / "slides.pptx"
    _create_image_pptx(path, ["Slide one", "Slide two"])
    ocr = FakeOCR(["OCR slide one", "OCR slide two"])

    result = PPTXProcessor(ocr).process(path)

    assert result.extraction_method == "ocr"
    assert [(section.source_type, section.source_index) for section in result.sections] == [
        ("slide", 1),
        ("slide", 2),
    ]
    assert [section.text for section in result.sections] == ["OCR slide one", "OCR slide two"]
    assert all(section.extraction_method == "ocr" for section in result.sections)
    assert ocr.calls == 2


def test_text_pptx_does_not_call_ocr(tmp_path: Path) -> None:
    path = tmp_path / "text-slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Text slide"
    presentation.save(path)
    ocr = FakeOCR(["should not be used"])

    result = PPTXProcessor(ocr).process(path)

    assert result.extraction_method == "text"
    assert result.sections[0].extraction_method == "text"
    assert ocr.calls == 0


def test_empty_ocr_result_is_clear(tmp_path: Path) -> None:
    path = tmp_path / "empty-slides.pptx"
    _create_image_pptx(path, ["Unreadable"])

    with pytest.raises(EmptyDocumentError, match="no extractable text"):
        PPTXProcessor(FakeOCR([""])).process(path)


def test_ocr_failure_is_reported(tmp_path: Path) -> None:
    path = tmp_path / "broken.pdf"
    _create_image_pdf(path)

    with pytest.raises(EmptyDocumentError, match="OCR is unavailable"):
        PDFProcessor(FakeOCR(error=OCRProcessingError("OCR engine failed"))).process(path)


def test_docx_processor_remains_text_based(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "lecture.docx"
    document = Document()
    document.add_paragraph("DOCX text")
    document.save(path)

    result = DOCXProcessor().process(path)

    assert result.extraction_method == "text"
    assert result.sections[0].extraction_method == "text"
