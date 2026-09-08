"""PPTX text extraction, preserving slide traceability."""

from pathlib import Path
from io import BytesIO

from PIL import Image
from pptx import Presentation
from pptx.exc import PackageNotFoundError

from app.processors.models import (
    Document,
    DocumentNotFoundError,
    DocumentProcessor,
    DocumentReadError,
    EmptyDocumentError,
    Section,
)
from app.processors.ocr import LocalOCRProvider, OCRProcessingError, OCRProvider


class PPTXProcessor(DocumentProcessor):
    """Extract slide text with OCR fallback for image-only slides."""

    supported_extension = ".pptx"

    def __init__(self, ocr_provider: OCRProvider | None = None) -> None:
        self._ocr_provider = ocr_provider

    def process(self, file_path: Path) -> Document:
        if not file_path.exists() or not file_path.is_file():
            raise DocumentNotFoundError(f"Document does not exist: {file_path}")
        sections: list[Section] = []
        used_ocr = False
        try:
            presentation = Presentation(file_path)
            for slide_number, slide in enumerate(presentation.slides, start=1):
                texts = []
                for shape in slide.shapes:
                    texts.extend(_shape_text(shape))
                text = "\n".join(value for value in texts if value).strip()
                extraction_method = "text"
                if not text:
                    images = _shape_images(slide.shapes)
                    if images:
                        text = "\n".join(self._ocr(image) for image in images).strip()
                        used_ocr = True
                        extraction_method = "ocr"
                if text:
                    sections.append(
                        Section(
                            source_index=slide_number,
                            source_type="slide",
                            text=text,
                            extraction_method=extraction_method,
                        )
                    )
                elif _shape_images(slide.shapes):
                    raise EmptyDocumentError(
                        f"Slide {slide_number} contains no extractable text and OCR could not detect readable content."
                    )
        except (OSError, PackageNotFoundError, ValueError) as error:
            raise DocumentReadError(
                f"Could not read PPTX document: {error}"
            ) from error

        if not sections:
            raise EmptyDocumentError("The PPTX presentation contains no extractable text or images.")
        return Document(
            file_path.name,
            self.supported_extension,
            tuple(sections),
            extraction_method="ocr" if used_ocr else "text",
        )

    def _ocr(self, image: Image.Image) -> str:
        provider = self._ocr_provider or LocalOCRProvider()
        try:
            return provider.extract_text(image).strip()
        except OCRProcessingError:
            raise
        except Exception as error:
            raise DocumentReadError(f"PPTX OCR failed: {error}") from error


def _shape_text(shape: object) -> list[str]:
    """Extract text from text boxes and table cells where available."""
    values: list[str] = []
    if getattr(shape, "has_text_frame", False):
        text = shape.text.strip()
        if text:
            values.append(text)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    values.append(text)
    return values


def _shape_images(shapes: object) -> list[Image.Image]:
    images: list[Image.Image] = []
    for shape in shapes:
        if getattr(shape, "shape_type", None) == 13:
            try:
                images.append(Image.open(BytesIO(shape.image.blob)))
            except (AttributeError, OSError) as error:
                raise DocumentReadError(f"Could not read a slide image: {error}") from error
    return images
