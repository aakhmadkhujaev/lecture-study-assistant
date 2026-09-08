"""DOCX text extraction, preserving paragraph order."""

from pathlib import Path

from docx import Document as WordDocument
from docx.opc.exceptions import PackageNotFoundError

from app.processors.models import (
    Document,
    DocumentNotFoundError,
    DocumentProcessor,
    DocumentReadError,
    EmptyDocumentError,
    Section,
)


class DOCXProcessor(DocumentProcessor):
    """Extract non-empty paragraphs from a DOCX in document order."""

    supported_extension = ".docx"

    def process(self, file_path: Path) -> Document:
        if not file_path.exists() or not file_path.is_file():
            raise DocumentNotFoundError(f"Document does not exist: {file_path}")
        try:
            word_document = WordDocument(file_path)
            sections = tuple(
                Section(
                    source_index=paragraph_index,
                    source_type=_paragraph_source_type(paragraph.style.name),
                    text=paragraph.text.strip(),
                )
                for paragraph_index, paragraph in enumerate(
                    word_document.paragraphs,
                    start=1,
                )
                if paragraph.text.strip()
            )
        except (OSError, PackageNotFoundError, ValueError) as error:
            raise DocumentReadError(
                f"Could not read DOCX document: {error}"
            ) from error

        if not sections:
            raise EmptyDocumentError("The DOCX document contains no extractable text.")
        return Document(file_path.name, self.supported_extension, sections)


def _paragraph_source_type(style_name: str) -> str:
    """Preserve heading information without interpreting document meaning."""
    return "heading" if style_name.lower().startswith("heading") else "paragraph"
